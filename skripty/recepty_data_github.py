# -*- coding: utf-8 -*-
"""
recepty_data_github.py — stiahne výživové tabuľky (USDA FoodData Central, Ciqual) a výrobky z Open Food Facts
so Slovenskom ako krajinou predaja a pripraví z nich kompaktné súbory pre vetvu `data-zdroje` repa
github.com/mrlukas161/Recepty. Cloud agent appky Recepty má sieť len na GitHub a npm, preto dáta podáva PC.

PREČO: Lukáš 16.10.2026 15:34 – „Sieť nechaj tak, stiahni výživové tabuľky a slovenské výrobky cez PC a podaj
ich cloudu cez GitHub." (konverzácia 191, otázka #414).

Príkazy (py -3 -I recepty_data_github.py <príkaz> <priečinok> …):
  usda <surove> <vystup>      stiahne SR Legacy + Foundation Foods (CSV zip), vytvorí <vystup>/usda/usda_vyziva.csv
  ciqual <surove> <vystup>    stiahne Ciqual 2025 (XLSX + alim XML), vytvorí <vystup>/ciqual/ciqual_vyziva.csv
  off-kody <vystup>           zoznam kódov výrobkov s krajinou Slovensko (search.openfoodfacts.org) → <vystup>/off/kody.txt
  off-produkty <vystup>       podrobnosti výrobkov cez OFF API v2 (najviac 85 dopytov/min, pokračuje po prerušení)
                              → <vystup>/off/sk-vyrobky.jsonl, na konci .jsonl.gz
Nič nemaže mimo zadaných priečinkov; stiahnuté súbory číta len ako dáta.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import os
import sys
import time
import zipfile
from datetime import datetime, timezone
from xml.etree import ElementTree as ET

import requests

UA = "Recepty/1.0 (github.com/mrlukas161/Recepty)"
USDA = {
    "sr_legacy": "https://fdc.nal.usda.gov/fdc-datasets/FoodData_Central_sr_legacy_food_csv_2018-04.zip",
    "foundation": "https://fdc.nal.usda.gov/fdc-datasets/FoodData_Central_foundation_food_csv_2025-12-18.zip",
}
# živiny podľa názvu v nutrient.csv (overí sa pri behu a vypíše)
USDA_ZIVINY = {
    "kcal": [("Energy", "KCAL")],
    "kcal_atwater_vseob": [("Energy (Atwater General Factors)", "KCAL")],
    "kcal_atwater_spec": [("Energy (Atwater Specific Factors)", "KCAL")],
    "bielkoviny_g": [("Protein", "G")],
    "tuky_g": [("Total lipid (fat)", "G")],
    "nasytene_g": [("Fatty acids, total saturated", "G")],
    "sacharidy_g": [("Carbohydrate, by difference", "G")],
    "cukry_g": [("Sugars, total including NLEA", "G"), ("Total Sugars", "G"), ("Sugars, Total", "G")],
    "vlaknina_g": [("Fiber, total dietary", "G")],
    "sodik_mg": [("Sodium, Na", "MG")],
}
CIQUAL_DOI = "doi:10.57745/RDMHWY"
CIQUAL_API = "https://entrepot.recherche.data.gouv.fr/api"
CIQUAL_SUBORY = ("Table Ciqual 2025_FR_2025_11_03.xlsx", "alim_2025_11_03.xml")
# stĺpce Ciqual podľa začiatku názvu hlavičky (overí sa pri behu a vypíše)
CIQUAL_STLPCE = {
    "kcal": "Energie, Règlement UE N° 1169/2011 (kcal/100 g)",
    "bielkoviny_g": "Protéines, N x facteur de Jones (g/100 g)",
    "sacharidy_g": "Glucides (g/100 g)",
    "tuky_g": "Lipides (g/100 g)",
    "cukry_g": "Sucres (g/100 g)",
    "vlaknina_g": "Fibres alimentaires (g/100 g)",
    "nasytene_g": "AG saturés (g/100 g)",
    "sol_g": "Sel chlorure de sodium (g/100 g)",
    "sodik_mg": "Sodium (mg/100 g)",
}
OFF_POLIA = ("code,product_name,product_name_sk,product_name_cs,product_name_en,brands,lang,quantity,"
             "ingredients_text,ingredients_text_sk,ingredients_text_cs,ingredients_text_en,ingredients_tags,"
             "additives_tags,ingredients_analysis_tags,allergens_tags,traces_tags,nutriments,nutriscore_grade,"
             "nova_group,categories_tags,last_modified_t")
OFF_ZIVINY = ("energy-kcal_100g", "fat_100g", "saturated-fat_100g", "carbohydrates_100g", "sugars_100g",
              "fiber_100g", "proteins_100g", "salt_100g", "sodium_100g")


def teraz() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def sha256(cesta: str) -> str:
    h = hashlib.sha256()
    with open(cesta, "rb") as f:
        for kus in iter(lambda: f.read(1 << 20), b""):
            h.update(kus)
    return h.hexdigest()


def stiahni(url: str, ciel: str) -> dict:
    os.makedirs(os.path.dirname(ciel), exist_ok=True)
    with requests.get(url, headers={"User-Agent": UA}, stream=True, timeout=120) as r:
        r.raise_for_status()
        with open(ciel, "wb") as f:
            for kus in r.iter_content(1 << 16):
                f.write(kus)
    return {"url": url, "subor": os.path.basename(ciel), "bajty": os.path.getsize(ciel), "sha256": sha256(ciel),
            "stiahnute": teraz()}


def zapis_zdroj(priecinok: str, data: dict) -> None:
    with open(os.path.join(priecinok, "zdroj.json"), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)


# ---------------------------------------------------------------- USDA
def usda(surove: str, vystup: str) -> None:
    out_dir = os.path.join(vystup, "usda")
    os.makedirs(out_dir, exist_ok=True)
    zdroje, riadky = [], []
    for typ, url in USDA.items():
        cesta = os.path.join(surove, os.path.basename(url))
        info = stiahni(url, cesta)
        zdroje.append({"typ": typ, **info})
        with zipfile.ZipFile(cesta) as z:
            mena = {os.path.basename(n): n for n in z.namelist() if n.endswith(".csv")}

            def citaj(meno):
                with z.open(mena[meno]) as f:
                    return list(csv.DictReader(io.TextIOWrapper(f, encoding="utf-8")))

            ziviny = citaj("nutrient.csv")
            id_pre = {}
            for kluc, moznosti in USDA_ZIVINY.items():
                for nazov, jednotka in moznosti:
                    n = next((x for x in ziviny if x["name"].strip() == nazov and x["unit_name"].upper() == jednotka), None)
                    if n:
                        id_pre[kluc] = n["id"]
                        print(f"  {typ}: {kluc} = id {n['id']} ({n['name']}, {n['unit_name']})")
                        break
                else:
                    print(f"  {typ}: {kluc} – v nutrient.csv sa nenašla")
            kat = {}
            if "food_category.csv" in mena:
                kat = {x["id"]: x["description"] for x in citaj("food_category.csv")}
            jedla = {x["fdc_id"]: x for x in citaj("food.csv")}
            if typ == "foundation":  # Foundation obsahuje aj vzorky a podvzorky; berieme len foundation_food
                jedla = {k: v for k, v in jedla.items() if v.get("data_type") == "foundation_food"}
            hodnoty: dict[str, dict] = {}
            hladane = {v: k for k, v in id_pre.items()}
            with z.open(mena["food_nutrient.csv"]) as f:
                for x in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8")):
                    if x["fdc_id"] in jedla and x["nutrient_id"] in hladane:
                        hodnoty.setdefault(x["fdc_id"], {})[hladane[x["nutrient_id"]]] = x["amount"]
            for fid, j in jedla.items():
                h = hodnoty.get(fid, {})
                sol = ""
                if h.get("sodik_mg") not in (None, ""):
                    sol = f"{float(h['sodik_mg']) * 2.5 / 1000:.3f}"
                riadky.append({"fdc_id": fid, "typ": typ, "popis": j.get("description", ""),
                               "kategoria": kat.get(j.get("food_category_id", ""), ""), **h, "sol_g_vypocet": sol})
    stlpce = ["fdc_id", "typ", "popis", "kategoria", *USDA_ZIVINY.keys(), "sol_g_vypocet"]
    cesta_csv = os.path.join(out_dir, "usda_vyziva.csv")
    with open(cesta_csv, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=stlpce)
        w.writeheader()
        for r in riadky:
            w.writerow({k: r.get(k, "") for k in stlpce})
    zapis_zdroj(out_dir, {"zdroj": "USDA FoodData Central (U.S. Department of Agriculture, Agricultural Research Service)",
                          "licencia": "CC0 1.0 (verejné vlastníctvo); citovať ako U.S. Department of Agriculture, FoodData Central",
                          "stranka": "https://fdc.nal.usda.gov/download-datasets", "subory": zdroje,
                          "vystup": {"subor": "usda_vyziva.csv", "riadkov": len(riadky), "sha256": sha256(cesta_csv)},
                          "poznamka": "Hodnoty na 100 g jedlej časti. sol_g_vypocet = sodík (mg) × 2,5 / 1000."})
    print(f"USDA: {len(riadky)} potravín → {cesta_csv}")


# ---------------------------------------------------------------- Ciqual
def ciqual(surove: str, vystup: str) -> None:
    out_dir = os.path.join(vystup, "ciqual")
    os.makedirs(out_dir, exist_ok=True)
    meta = requests.get(f"{CIQUAL_API}/datasets/:persistentId/", params={"persistentId": CIQUAL_DOI},
                        headers={"User-Agent": UA}, timeout=60).json()["data"]["latestVersion"]
    subory = {f["dataFile"]["filename"]: f["dataFile"]["id"] for f in meta["files"]}
    zdroje = []
    for meno in CIQUAL_SUBORY:
        cesta = os.path.join(surove, meno)
        zdroje.append(stiahni(f"{CIQUAL_API}/access/datafile/{subory[meno]}", cesta))
    # anglické názvy z alim XML
    en = {}
    for el in ET.parse(os.path.join(surove, CIQUAL_SUBORY[1])).getroot():
        kod = (el.findtext("alim_code") or "").strip()
        if kod:
            en[kod] = (el.findtext("alim_nom_eng") or "").strip()
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(surove, CIQUAL_SUBORY[0]), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    riadky_x = ws.iter_rows(values_only=True)
    hlavicka = [str(x or "").strip() for x in next(riadky_x)]
    # hlavičky v XLSX majú zalomenia riadkov namiesto medzier a „/“ → porovnávame len písmená a číslice
    norm = lambda s: "".join(ch for ch in s.lower() if ch.isalnum())
    hl_norm = [norm(h) for h in hlavicka]
    idx = {}
    for kluc, nazov in CIQUAL_STLPCE.items():
        i = next((n for n, h in enumerate(hl_norm) if h == norm(nazov)), None)
        print(f"  ciqual: {kluc} = stĺpec {i} ({' '.join(hlavicka[i].split()) if i is not None else 'NENAŠIEL SA'})")
        if i is not None:
            idx[kluc] = i
    zakl = {k: hlavicka.index(k) for k in ("alim_code", "alim_nom_fr", "alim_grp_nom_fr", "alim_ssgrp_nom_fr")
            if k in hlavicka}
    stlpce = ["alim_code", "nazov_fr", "nazov_en", "skupina", "podskupina", *CIQUAL_STLPCE.keys()]
    cesta_csv = os.path.join(out_dir, "ciqual_vyziva.csv")
    n = 0
    with open(cesta_csv, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=stlpce)
        w.writeheader()
        for r in riadky_x:
            kod = str(r[zakl["alim_code"]] or "").strip()
            if not kod:
                continue
            text = lambda v: " ".join(str(v).split()) if v is not None else ""
            zaz = {"alim_code": kod, "nazov_fr": text(r[zakl["alim_nom_fr"]]), "nazov_en": en.get(kod, ""),
                   "skupina": text(r[zakl["alim_grp_nom_fr"]]) if "alim_grp_nom_fr" in zakl else "",
                   "podskupina": text(r[zakl["alim_ssgrp_nom_fr"]]) if "alim_ssgrp_nom_fr" in zakl else ""}
            for kluc, i in idx.items():
                v = r[i]
                zaz[kluc] = "" if v is None else str(v).strip()
            w.writerow(zaz)
            n += 1
    zapis_zdroj(out_dir, {"zdroj": "Anses. 2025. Table de composition nutritionnelle des aliments Ciqual 2025",
                          "doi": "10.57745/RDMHWY", "licencia": "Licence Ouverte / Open Licence 2.0 (Etalab)",
                          "stranka": "https://entrepot.recherche.data.gouv.fr/dataset.xhtml?persistentId=doi:10.57745/RDMHWY",
                          "verzia": f"{meta.get('versionNumber')}.{meta.get('versionMinorNumber')} z {meta.get('releaseTime')}",
                          "subory": zdroje, "vystup": {"subor": "ciqual_vyziva.csv", "riadkov": n, "sha256": sha256(cesta_csv)},
                          "poznamka": "Hodnoty ako v tabuľke (desatinná čiarka, „< 0,5“, „traces“, „-“ ponechané); na 100 g."})
    print(f"Ciqual: {n} potravín → {cesta_csv}")


# ---------------------------------------------------------------- Open Food Facts
def off_kody(vystup: str) -> None:
    out_dir = os.path.join(vystup, "off")
    os.makedirs(out_dir, exist_ok=True)
    kody, strana, pocet_stran = [], 1, None
    while pocet_stran is None or strana <= pocet_stran:
        r = requests.get("https://search.openfoodfacts.org/search",
                         params={"q": 'countries_tags:"en:slovakia"', "page_size": 100, "page": strana, "fields": "code"},
                         headers={"User-Agent": UA}, timeout=60)
        r.raise_for_status()
        d = r.json()
        pocet_stran = d.get("page_count") or 0
        kody += [h["code"] for h in d.get("hits", []) if h.get("code")]
        print(f"  strana {strana}/{pocet_stran}: spolu {len(kody)}")
        strana += 1
        time.sleep(1.5)
    kody = list(dict.fromkeys(kody))
    with open(os.path.join(out_dir, "kody.txt"), "w", encoding="utf-8") as f:
        f.write("\n".join(kody) + "\n")
    print(f"OFF kódy: {len(kody)}")


def off_produkty(vystup: str) -> None:
    out_dir = os.path.join(vystup, "off")
    kody = [k.strip() for k in open(os.path.join(out_dir, "kody.txt"), encoding="utf-8") if k.strip()]
    cesta = os.path.join(out_dir, "sk-vyrobky.jsonl")
    hotove = set()
    if os.path.exists(cesta):
        for riadok in open(cesta, encoding="utf-8"):
            try:
                hotove.add(json.loads(riadok)["code"])
            except Exception:
                pass
    chyby, novych = 0, 0
    s = requests.Session()
    s.headers["User-Agent"] = UA
    with open(cesta, "a", encoding="utf-8") as f:
        for i, kod in enumerate(kody, 1):
            if kod in hotove:
                continue
            for pokus in range(5):
                try:
                    r = s.get(f"https://world.openfoodfacts.org/api/v2/product/{kod}.json",
                              params={"fields": OFF_POLIA}, timeout=60)
                    if r.status_code in (429, 502, 503, 504):
                        time.sleep(30 * (pokus + 1))
                        continue
                    d = r.json()
                    break
                except Exception:
                    time.sleep(10 * (pokus + 1))
            else:
                chyby += 1
                continue
            if d.get("status") == 1 and d.get("product"):
                p = d["product"]
                p["code"] = p.get("code") or kod
                nut = p.get("nutriments") or {}
                p["nutriments"] = {k: nut[k] for k in OFF_ZIVINY if k in nut}
                f.write(json.dumps(p, ensure_ascii=False) + "\n")
                f.flush()
                novych += 1
            if i % 100 == 0:
                print(f"  {i}/{len(kody)} (nových {novych}, chýb {chyby})", flush=True)
            time.sleep(0.7)  # najviac ~85 dopytov za minútu (limit OFF 100/min)
    with open(cesta, "rb") as f_in, gzip.open(cesta + ".gz", "wb", compresslevel=9) as f_out:
        f_out.write(f_in.read())
    n = sum(1 for _ in open(cesta, encoding="utf-8"))
    zapis_zdroj(out_dir, {"zdroj": "Open Food Facts (https://world.openfoodfacts.org)", "licencia":
                          "Open Database License (ODbL) 1.0; obsah databázy: Database Contents License; obrázky nesťahujeme",
                          "vyber": 'search.openfoodfacts.org q=countries_tags:"en:slovakia", podrobnosti cez /api/v2/product/{kód}.json',
                          "polia": OFF_POLIA, "ziviny": list(OFF_ZIVINY), "stiahnute": teraz(),
                          "vystup": {"subor": "sk-vyrobky.jsonl.gz", "vyrobkov": n, "sha256": sha256(cesta + ".gz")},
                          "chyby_stahovania": chyby})
    print(f"OFF: {n} výrobkov (chyby {chyby}) → {cesta}.gz")


if __name__ == "__main__":
    prikaz, *args = sys.argv[1:]
    {"usda": usda, "ciqual": ciqual, "off-kody": off_kody, "off-produkty": off_produkty}[prikaz](*args)
